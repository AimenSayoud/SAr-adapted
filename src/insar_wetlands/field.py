"""Field observations at Rzecin: water-table depth, laser surface level, UAV/LAI, plots.

This module is public; the data are not. It never contains or writes field values into the
repository: everything is read from the research hub's ``06_data/field`` (``RZECIN_FIELD_ROOT``)
and derived tables are written back there by the caller.

Conventions established from the first delivery (hub: ``06_data/field/raw/2026-09_abdallah_delivery/
SOURCE.md``):

* Logger stamps describe the hour **centred at stamp − 0.5 h UTC** (station air temperature vs
  ERA5 hourly t2m). ``load_wtd_hourly`` / ``load_laser`` return that centre as a UTC index.
* WTD in cm, negative below the surface. The hourly file is the primary hydrological series.
* Laser (Campbell SDMS40 snow-depth sensor at P6/CR): peat-surface position in cm relative to its
  calibration level; it also sees snow, hence ``snow_mask``.
* Merged UAV table: missing values coded −9999; sub-plots ``P1_1``…``P9_5`` are replicates of the
  main plot (P5_* = CL_*, P6_* = CR_*), which carries the WTD sensor.
"""
from __future__ import annotations

import os
from pathlib import Path

import numpy as np
import pandas as pd

UTC_CENTRE_OFFSET = pd.Timedelta(minutes=30)
DELIVERY = "raw/2026-09_abdallah_delivery"
WTD_COLUMNS = {"WTD_P1": "P1", "WTD_P2": "P2", "WTD_P3": "P3", "WTD_P4": "P4",
               "WTD_P5 (CL)": "P5", "WTD_P6 (CR)": "P6", "WTD_P7": "P7", "WTD_P8": "P8",
               "WTD_P9": "P9"}
PLOTS = list(WTD_COLUMNS.values())


def field_root(root: str | Path | None = None) -> Path:
    """The hub's ``06_data/field``: explicit argument, ``RZECIN_FIELD_ROOT``, or the hub that
    contains this repository (``05_code/SAr-adapted`` → ``../../06_data/field``)."""
    if root is not None:
        return Path(root)
    if os.environ.get("RZECIN_FIELD_ROOT"):
        return Path(os.environ["RZECIN_FIELD_ROOT"])
    return Path(__file__).resolve().parents[4] / "06_data" / "field"


def base_plot(name: str) -> str:
    """Main (WTD) plot of a plot or replicate label: ``P5_2``→``P5``, ``CR_1``→``P6``, ``P8_2023``→``P8``."""
    name = str(name)
    if name.startswith("CL"):
        return "P5"
    if name.startswith("CR"):
        return "P6"
    return name.split("_")[0]


# ------------------------------------------------------------------------------ readers

def load_wtd_hourly(path: str | Path | None = None) -> pd.DataFrame:
    """Hourly WTD (cm) for P1–P9 plus the raw CL/CR levels and meteo, indexed by UTC centre time.

    Columns: ``P1``…``P9`` (WTD), ``CL_raw``, ``CR_raw``, and the meteo columns as delivered.
    """
    path = Path(path) if path else field_root() / DELIVERY / "WTD_hourly_2020-2024_9plots_filled_meteo.csv"
    df = pd.read_csv(path, parse_dates=["TIMESTAMP"])
    df.index = pd.DatetimeIndex(df.pop("TIMESTAMP") - UTC_CENTRE_OFFSET, name="time_utc").tz_localize("UTC")
    df = df.rename(columns={**WTD_COLUMNS, "CL": "CL_raw", "CR": "CR_raw"})
    if df.index.has_duplicates or not df.index.is_monotonic_increasing:
        raise ValueError(f"{path.name}: time index is not strictly increasing")
    return df


def load_laser(path: str | Path | None = None) -> pd.DataFrame:
    """Laser surface level at P6/CR (cm), the surface series used for the WTD correction, and
    the corrected WTD; UTC centre-time index. Rows without a timestamp are dropped."""
    path = Path(path) if path else field_root() / DELIVERY / "Laser_Sensor.xlsx"
    raw = pd.read_excel(path, sheet_name=0)
    raw.columns = ["time", "surface_cm", "surface_ref_cm", "wtd_corrected_cm"]
    raw = raw[raw["time"].notna()].copy()
    idx = pd.DatetimeIndex(pd.to_datetime(raw.pop("time")) - UTC_CENTRE_OFFSET, name="time_utc").tz_localize("UTC")
    out = raw.apply(pd.to_numeric, errors="coerce")
    out.index = idx
    return out


def snow_mask(air_temp_c: pd.Series, index: pd.DatetimeIndex, hours: int = 72) -> pd.Series:
    """True where the laser may be seeing snow: any sub-zero hourly air temperature in the
    ``hours`` before (and including) each time. Aligned to ``index``. Times outside the air
    record (e.g. laser hours after the meteo series ends) are True: unknown is not snow-free."""
    cold = (air_temp_c < 0).astype(float).rolling(f"{hours}h", min_periods=1).max() > 0
    out = cold.reindex(index, method="ffill").fillna(True).astype(bool)
    outside = (index < air_temp_c.index[0]) | (index > air_temp_c.index[-1])
    return out | pd.Series(outside, index=index)


FLAG_FILL = "FFFFFF00"   # the field team marks doubtful cells in solid yellow


def flagged_cells(path: str | Path, sheet: str = "Sheet1", fill: str = FLAG_FILL) -> list[tuple[int, str]]:
    """(data row index, column name) of every cell filled with ``fill`` — the field team's
    "exclude" marking in the merged table (row 0 = first row under the header)."""
    import openpyxl
    ws = openpyxl.load_workbook(path, read_only=False)[sheet]
    header = [c.value for c in ws[1]]
    out = []
    for r, row in enumerate(ws.iter_rows(min_row=2)):
        for c in row:
            f = c.fill
            if f is not None and f.fill_type == "solid" and f.fgColor is not None and f.fgColor.type == "rgb" \
                    and f.fgColor.rgb == fill and c.column - 1 < len(header):
                out.append((r, header[c.column - 1]))
    return out


def load_uav_table(path: str | Path | None = None, mask_flagged: bool = True) -> pd.DataFrame:
    """Merged UAV/LAI/meteo/WTD table with −9999 → NaN, parsed dates and a ``base_plot`` column.
    ``mask_flagged``: cells the field team marked yellow are set to NaN (their count is in
    ``df.attrs["n_flagged"]``)."""
    path = Path(path) if path else field_root() / DELIVERY / "DataSet_All_RS_LAI_merged_with_WTD_Meteo.xlsx"
    df = pd.read_excel(path, sheet_name="Sheet1").replace(-9999, np.nan)
    flags = flagged_cells(path) if mask_flagged else []
    for r, col in flags:
        if col in df.columns and col not in ("Date", "Plot"):
            df.loc[df.index[r], col] = np.nan
    df.attrs["n_flagged"] = len(flags)
    df["Date"] = pd.to_datetime(df["Date"])
    df["base_plot"] = df["Plot"].map(base_plot)
    return df


def load_plots(path: str | Path | None = None, crs: int = 32633):
    """Plot polygons with ``plot``, ``base_plot`` and centroid ``x``/``y`` in ``crs``."""
    import geopandas as gpd
    path = Path(path) if path else field_root() / DELIVERY / "plots" / "Plots_PlanetScope.shp"
    g = gpd.read_file(path).to_crs(crs).rename(columns={"Plot": "plot", "Treatment": "treatment"})
    g["base_plot"] = g["plot"].map(base_plot)
    c = g.geometry.centroid
    g["x"], g["y"] = c.x, c.y
    return g


# ------------------------------------------------------------------------------ Sentinel-1 matching

def wtd_at_times(wtd: pd.DataFrame, times_utc, plots=PLOTS) -> pd.DataFrame:
    """The WTD variables the supervisor asked for, for each plot at each acquisition time.

    ``wtd_at`` (linear interpolation at the time), ``mean_24h`` (24 h centred), ``mean_prev3d`` /
    ``mean_prev7d`` (windows ending at the time), ``change_3d`` / ``change_7d`` (value at the time
    minus value 3 / 7 days before). One row per (time, plot).
    """
    times = pd.DatetimeIndex(pd.to_datetime(times_utc, utc=True))
    rows = []
    for p in plots:
        s = wtd[p].astype(float)
        for t in times:
            v = _interp_at(s, t)
            rows.append({
                "time_utc": t, "plot": p, "wtd_at": v,
                "mean_24h": s[t - pd.Timedelta(hours=12): t + pd.Timedelta(hours=12)].mean(),
                "mean_prev3d": s[t - pd.Timedelta(days=3): t].mean(),
                "mean_prev7d": s[t - pd.Timedelta(days=7): t].mean(),
                "change_3d": v - _interp_at(s, t - pd.Timedelta(days=3)),
                "change_7d": v - _interp_at(s, t - pd.Timedelta(days=7)),
            })
    return pd.DataFrame(rows)


def _interp_at(s: pd.Series, t: pd.Timestamp) -> float:
    """Linear interpolation of an hourly series at time t; NaN outside the record. Works in
    seconds from the first sample, whatever the index's time unit (pandas ≥ 3 is not ns-only)."""
    if not (s.index[0] <= t <= s.index[-1]):
        return np.nan
    x = (s.index - s.index[0]).total_seconds().to_numpy()
    return float(np.interp((t - s.index[0]).total_seconds(), x, s.to_numpy()))


def surface_wetness_at(meteo: pd.DataFrame, times_utc, rh_wet: float = 95.0,
                       rain_hours: int = 3) -> pd.DataFrame:
    """Surface state at each overpass from the station meteo (``load_wtd_hourly`` columns
    ``RH_2m``, ``Air_2m``, ``Rain_mm_Tot``).

    ``wet``: relative humidity ≥ ``rh_wet`` % at the time (dew or wet canopy likely) or rain in
    the ``rain_hours`` before it. ``frozen``: air temperature ≤ 0 °C at the time — a different
    dielectric state, kept apart from wetness. One row per time.
    """
    times = pd.DatetimeIndex(pd.to_datetime(times_utc, utc=True))
    rain = meteo["Rain_mm_Tot"].astype(float)
    rows = []
    for t in times:
        rh = _interp_at(meteo["RH_2m"].astype(float), t)
        air = _interp_at(meteo["Air_2m"].astype(float), t)
        r = float(rain[(rain.index > t - pd.Timedelta(hours=rain_hours)) & (rain.index <= t)].sum())
        rows.append({"time_utc": t, "rh": rh, "air_c": air, "rain_prev_mm": r,
                     "wet": bool((rh >= rh_wet) or (r > 0)), "frozen": bool(air <= 0)})
    return pd.DataFrame(rows)


def window_stats(values: np.ndarray, valid: np.ndarray | None = None) -> dict:
    """Statistics of the pixels in an extraction window (NaN = invalid)."""
    v = np.asarray(values, float).ravel()
    if valid is not None:
        v = np.where(np.asarray(valid, bool).ravel(), v, np.nan)
    v = v[np.isfinite(v)]
    if not v.size:
        return {"n_valid": 0, "median": np.nan, "mean": np.nan, "sd": np.nan, "iqr": np.nan}
    q1, q3 = np.percentile(v, [25, 75])
    return {"n_valid": int(v.size), "median": float(np.median(v)), "mean": float(v.mean()),
            "sd": float(v.std(ddof=1)) if v.size > 1 else np.nan, "iqr": float(q3 - q1)}


def censored_flag(s: pd.Series, floor_quantile: float = 0.02, tol_cm: float = 1.5,
                  min_days: float = 5.0) -> pd.Series:
    """True where an hourly WTD series probably sits at the bottom of the well / sensor: within
    ``tol_cm`` of the plot's low floor (its ``floor_quantile`` quantile) for at least ``min_days``
    in a row. Such stretches (e.g. P1, P4–P6, P9 in Aug–Dec 2022) bound the water table from
    above; they do not measure it. Brief dips to the same depth are not flagged. Heuristic —
    to be replaced by the logger's own flag if the field team has one."""
    near = s <= s.quantile(floor_quantile) + tol_cm
    run_id = (near != near.shift()).cumsum()
    run_len = near.groupby(run_id).transform("size")
    step_h = (s.index[1] - s.index[0]).total_seconds() / 3600 if len(s) > 1 else 1.0
    return (near & (run_len * step_h >= min_days * 24)).astype(bool)
