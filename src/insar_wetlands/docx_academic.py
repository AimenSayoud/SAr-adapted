"""
Academic DOCX Post-Processor for Scientific Manuscripts.

Converts standard Pandoc DOCX output into publication-grade camera-ready
formatting adhering to journal standards (Elsevier, IEEE, Nature):
1. Applies Booktabs table formatting (1.5pt top/bottom, 0.75pt midrule, 0 vertical).
2. Centers tables, adds cell padding, sets repeated headers and prevents row splitting.
3. Cleans heading styles to pure black (eliminates blue theme accents and markers).
4. Standardizes table and figure captions (bold identifiers, keep-with-next).
5. Centers figures and maintains figure-caption pairing.
"""

import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Pt, RGBColor

FONT_NAME = "Times New Roman"
COLOR_BLACK = RGBColor(0, 0, 0)


def apply_booktabs_to_table(table) -> None:
    """Applies publication-grade Booktabs styling to a python-docx Table."""
    tblPr = table._tbl.tblPr

    # 1. Center the table horizontally
    existing_jc = tblPr.find(qn("w:jc"))
    if existing_jc is not None:
        tblPr.remove(existing_jc)
    tblPr.append(parse_xml(f'<w:jc {nsdecls("w")} w:val="center"/>'))

    # 2. Table-Level Borders (Top 1.5pt, Bottom 1.5pt, Zero vertical/interior)
    existing_tblBorders = tblPr.find(qn("w:tblBorders"))
    if existing_tblBorders is not None:
        tblPr.remove(existing_tblBorders)

    tblBorders = parse_xml(f"""
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="12" w:space="0" w:color="000000"/>
            <w:left w:val="none" w:sz="0" w:space="0" w:color="auto"/>
            <w:bottom w:val="single" w:sz="12" w:space="0" w:color="000000"/>
            <w:right w:val="none" w:sz="0" w:space="0" w:color="auto"/>
            <w:insideH w:val="none" w:sz="0" w:space="0" w:color="auto"/>
            <w:insideV w:val="none" w:sz="0" w:space="0" w:color="auto"/>
        </w:tblBorders>
    """)
    tblPr.append(tblBorders)

    # 3. Cell Margins / Padding (120 dxa = 6pt top/bottom; 160 dxa = 8pt left/right)
    existing_cellMar = tblPr.find(qn("w:tblCellMar"))
    if existing_cellMar is not None:
        tblPr.remove(existing_cellMar)

    cellMar = parse_xml(f"""
        <w:tblCellMar {nsdecls("w")}>
            <w:top w:w="120" w:type="dxa"/>
            <w:bottom w:w="120" w:type="dxa"/>
            <w:left w:w="160" w:type="dxa"/>
            <w:right w:w="160" w:type="dxa"/>
        </w:tblCellMar>
    """)
    tblPr.append(cellMar)

    n_rows = len(table.rows)
    if n_rows == 0:
        return

    # 4. Header Row (Row 0): Repeated header, prevent split, 0.75pt midrule
    trPr0 = table.rows[0]._tr.get_or_add_trPr()
    if trPr0.find(qn("w:tblHeader")) is None:
        trPr0.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    if trPr0.find(qn("w:cantSplit")) is None:
        trPr0.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

    for cell in table.rows[0].cells:
        tcPr = cell._tc.get_or_add_tcPr()
        existing_tcBorders = tcPr.find(qn("w:tcBorders"))
        if existing_tcBorders is not None:
            tcPr.remove(existing_tcBorders)

        # 0.75 pt (sz="6") solid black line below header cells
        tcBorders = parse_xml(f"""
            <w:tcBorders {nsdecls("w")}>
                <w:bottom w:val="single" w:sz="6" w:space="0" w:color="000000"/>
            </w:tcBorders>
        """)
        tcPr.append(tcBorders)

        for p in cell.paragraphs:
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            for r in p.runs:
                r.font.name = FONT_NAME
                r.font.size = Pt(10)
                r.font.bold = True
                r.font.color.rgb = COLOR_BLACK

    # 5. Body Rows (Rows 1 to N): Prevent row split across pages, clean borders
    for r_idx in range(1, n_rows):
        row = table.rows[r_idx]
        trPr = row._tr.get_or_add_trPr()
        if trPr.find(qn("w:cantSplit")) is None:
            trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            existing_tcBorders = tcPr.find(qn("w:tcBorders"))
            if existing_tcBorders is not None:
                tcPr.remove(existing_tcBorders)

            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                for r in p.runs:
                    r.font.name = FONT_NAME
                    r.font.size = Pt(9.5)
                    r.font.color.rgb = COLOR_BLACK


def polish_academic_docx(docx_in: str, docx_out: str) -> dict:
    """Polishes a Word manuscript DOCX to professional academic journal standards."""
    doc = Document(docx_in)

    n_tables = len(doc.tables)
    for table in doc.tables:
        apply_booktabs_to_table(table)

    n_headings = 0
    n_captions = 0
    n_figures = 0

    for p in doc.paragraphs:
        # 1. Clean Heading Typography & Pagination
        if p.style and p.style.name.startswith("Heading"):
            n_headings += 1
            p.paragraph_format.keep_with_next = True
            for r in p.runs:
                r.font.name = FONT_NAME
                r.font.color.rgb = COLOR_BLACK

        text = p.text.strip()

        # 2. Table Captions (Top of Table, Keep with next)
        if text.startswith(("Table ", "Tab. ")) or ("**Table " in text):
            n_captions += 1
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            for r in p.runs:
                r.font.name = FONT_NAME
                r.font.size = Pt(10)
                if any(k in r.text for k in ["Table ", "Tab. "]):
                    r.font.bold = True
                    r.font.color.rgb = COLOR_BLACK

        # 3. Figure Captions (Bottom of Figure)
        elif text.startswith(("Figure ", "Fig. ")) or ("**Figure " in text):
            n_captions += 1
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(12)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.name = FONT_NAME
                r.font.size = Pt(10)
                if any(k in r.text for k in ["Figure ", "Fig. "]):
                    r.font.bold = True
                    r.font.color.rgb = COLOR_BLACK

        # 4. Figure Images Alignment
        if p._p.xpath(".//a:blip") or p._p.xpath(".//w:drawing"):
            n_figures += 1
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True

    Path(docx_out).parent.mkdir(parents=True, exist_ok=True)
    doc.save(docx_out)

    return {
        "tables_polished": n_tables,
        "headings_cleaned": n_headings,
        "captions_styled": n_captions,
        "figures_centered": n_figures,
        "output": docx_out,
    }


if __name__ == "__main__":
    if len(sys.argv) < 3:
        print("Usage: python docx_academic.py <input.docx> <output.docx>")
        sys.exit(1)
    res = polish_academic_docx(sys.argv[1], sys.argv[2])
    print(res)
