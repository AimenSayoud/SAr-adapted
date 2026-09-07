"""Tests for the academic DOCX post-processing module."""

import tempfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from insar_wetlands.docx_academic import apply_booktabs_to_table, polish_academic_docx


def test_apply_booktabs_to_table():
    doc = Document()
    t = doc.add_table(rows=3, cols=2)
    for r_idx, row in enumerate(t.rows):
        for c_idx, cell in enumerate(row.cells):
            cell.text = f"cell_{r_idx}_{c_idx}"

    apply_booktabs_to_table(t)

    # 1. Verify centered alignment
    jc = t._tbl.tblPr.find(qn("w:jc"))
    assert jc is not None
    assert jc.attrib[qn("w:val")] == "center"

    # 2. Verify table borders (top and bottom sz=12, others none)
    tblBorders = t._tbl.tblPr.find(qn("w:tblBorders"))
    assert tblBorders is not None
    top = tblBorders.find(qn("w:top"))
    assert top is not None and top.attrib[qn("w:sz")] == "12"
    bottom = tblBorders.find(qn("w:bottom"))
    assert bottom is not None and bottom.attrib[qn("w:sz")] == "12"
    left = tblBorders.find(qn("w:left"))
    assert left is not None and left.attrib[qn("w:val")] == "none"

    # 3. Verify repeating header row and cantSplit
    tr0 = t.rows[0]._tr.trPr
    assert tr0.find(qn("w:tblHeader")) is not None
    assert tr0.find(qn("w:cantSplit")) is not None

    # 4. Verify header cell midrule (sz=6)
    tc0 = t.rows[0].cells[0]._tc.tcPr
    tcBorders = tc0.find(qn("w:tcBorders"))
    assert tcBorders is not None
    assert tcBorders.find(qn("w:bottom")).attrib[qn("w:sz")] == "6"

    # 5. Verify data row cantSplit
    tr1 = t.rows[1]._tr.trPr
    assert tr1.find(qn("w:cantSplit")) is not None


def test_polish_academic_docx_pipeline():
    with tempfile.TemporaryDirectory() as td:
        p_in = Path(td) / "in.docx"
        p_out = Path(td) / "out.docx"

        doc = Document()
        doc.add_heading("1. Introduction", level=1)
        doc.add_paragraph("Table 1. Experimental summary of zones.")
        t = doc.add_table(rows=2, cols=2)
        t.rows[0].cells[0].text = "Header 1"
        t.rows[0].cells[1].text = "Header 2"
        t.rows[1].cells[0].text = "123"
        t.rows[1].cells[1].text = "456"
        doc.add_paragraph("Figure 1. Overview of the wetland complex.")

        doc.save(p_in)

        res = polish_academic_docx(str(p_in), str(p_out))
        assert res["tables_polished"] == 1
        assert res["headings_cleaned"] >= 1
        assert res["captions_styled"] >= 2
        assert p_out.exists()

        # Reopen output document to verify styles
        out_doc = Document(p_out)
        assert len(out_doc.tables) == 1
        tblBorders = out_doc.tables[0]._tbl.tblPr.find(qn("w:tblBorders"))
        assert tblBorders is not None
        assert tblBorders.find(qn("w:top")).attrib[qn("w:sz")] == "12"
